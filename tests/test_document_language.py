"""Belge Dili (TR / EN) ve Başlıksız Ek Alan testleri."""

from __future__ import annotations

from datetime import date
from decimal import Decimal
from pathlib import Path
import sqlite3

import openpyxl
from PySide6.QtPdf import QPdfDocument
import pytest

from teklif_sablonu.domain.calculator import calculate_offer
from teklif_sablonu.domain.enums import (
    ColumnKey,
    CurrencyCode,
    DiscountType,
    DocumentLanguage,
    PageOrientation,
    QuantityUnit,
    SectionKey,
)
from teklif_sablonu.domain.models import Offer, OfferItem, OfferSection
from teklif_sablonu.services.excel_service import ExcelService
from teklif_sablonu.services.pdf_service import PdfService
from teklif_sablonu.storage.repository import SqliteOfferRepository


def test_document_language_enum_and_labels():
    assert DocumentLanguage.TR.value == "tr"
    assert DocumentLanguage.EN.value == "en"

    # Sütun başlıkları
    assert ColumnKey.DESCRIPTION.label(DocumentLanguage.TR) == "Ürün Açıklaması"
    assert ColumnKey.DESCRIPTION.label(DocumentLanguage.EN) == "Description"
    assert ColumnKey.UNIT_PRICE.label(DocumentLanguage.EN) == "Unit Price"

    # Miktar birimi
    assert QuantityUnit.PIECE.label(DocumentLanguage.TR) == "Adet"
    assert QuantityUnit.PIECE.label(DocumentLanguage.EN) == "Piece"
    assert QuantityUnit.LITRE.label(DocumentLanguage.EN) == "Liter"
    assert QuantityUnit.PIECE.get_column_header(DocumentLanguage.EN) == "Quantity (Piece)"

    # İskonto tipi
    assert DiscountType.RATE.label(DocumentLanguage.EN) == "Rate"
    assert DiscountType.AMOUNT.label(DocumentLanguage.EN) == "Amount"

    # Koşul başlıkları
    assert SectionKey.DELIVERY_ADDRESS.label(DocumentLanguage.EN) == "Delivery Address"
    assert SectionKey.PAYMENT_TERMS.label(DocumentLanguage.EN) == "Payment Terms"


def test_repository_save_and_retrieve_document_language():
    conn = sqlite3.connect(":memory:")
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON;")
    repo = SqliteOfferRepository(conn)

    offer = Offer(
        offer_number="20261021-001",
        offer_date=date(2026, 10, 21),
        company_id="endemiks_tarim",
        recipient_name="Global Corp LLC",
        document_language=DocumentLanguage.EN,
        items=[
            OfferItem(
                product_code="EN-01",
                description="Organic Fertilizer",
                quantity=Decimal("10"),
                unit_price=Decimal("150.00"),
            )
        ],
    )

    saved = repo.save(offer)
    assert saved.id is not None
    assert saved.document_language == DocumentLanguage.EN

    retrieved = repo.get_by_id(saved.id)
    assert retrieved is not None
    assert retrieved.document_language == DocumentLanguage.EN

    # Güncelleme testi
    retrieved.document_language = DocumentLanguage.TR
    updated = repo.save(retrieved)
    assert updated.document_language == DocumentLanguage.TR


def test_pdf_flowables_and_generation_english(tmp_path: Path):
    pdf_path = tmp_path / "test_en.pdf"
    service = PdfService()

    offer = Offer(
        offer_number="20261021-002",
        offer_date=date(2026, 10, 21),
        company_id="endemiks_tarim",
        recipient_name="Acme International",
        currency_code=CurrencyCode.USD,
        quantity_unit=QuantityUnit.PIECE,
        vat_enabled=True,
        vat_rate=Decimal("20.00"),
        discount_enabled=True,
        document_language=DocumentLanguage.EN,
        selected_columns=[
            ColumnKey.DESCRIPTION,
            ColumnKey.QUANTITY,
            ColumnKey.UNIT_PRICE,
            ColumnKey.DISCOUNT_TYPE,
            ColumnKey.TOTAL_AMOUNT,
        ],
        items=[
            OfferItem(
                product_code="US-01",
                description="Export Grade Seed",
                quantity=Decimal("100"),
                unit_price=Decimal("25.00"),
                discount_type=DiscountType.RATE,
                discount_value=Decimal("10.00"),
            )
        ],
        sections=[
            OfferSection(
                section_key=SectionKey.DELIVERY_ADDRESS,
                text="FOB Istanbul Port within 10 days.",
                enabled=True,
            ),
            OfferSection(
                section_key=SectionKey.ADDITIONAL_NOTES,
                text="<p>Special export packaging and pallets included.</p>",
                enabled=True,
                rich_text=True,
            ),
        ],
    )

    # 1. Header flowables testi
    header_table = service._build_header(offer, 500)
    header_cells_text = str(header_table._cellvalues)
    assert "OFFER NO:" in header_cells_text
    assert "DATE:" in header_cells_text
    assert "TEKLİF NO:" not in header_cells_text
    assert "Headquarters and Factory:" in header_cells_text
    assert "Merkez Ofis ve Fabrika:" not in header_cells_text

    # 2. Items table flowables testi
    valid_items = offer.get_valid_items()
    calcs, totals = calculate_offer(offer)
    visible_cols = service.filter_visible_columns(offer, valid_items)
    items_table = service._build_items_table(offer, valid_items, calcs, visible_cols, 500)
    items_text = str(items_table._cellvalues)
    assert "DESCRIPTION" in items_text
    assert "QUANTITY (PIECE)" in items_text
    assert "UNIT PRICE" in items_text
    assert "Rate" in items_text

    # 3. Totals flowables testi
    totals_table = service._build_totals_table(offer, totals, 500)
    totals_text = str(totals_table._cellvalues)
    assert "Subtotal:" in totals_text
    assert "Total Discount:" in totals_text
    assert "Total VAT" in totals_text
    assert "GRAND TOTAL (INCL. VAT):" in totals_text
    assert "Ara Toplam:" not in totals_text

    # 4. Conditions flowables testi: Ek Alan başlıksız olmalı
    conds_block = service._build_conditions(offer, 500)
    assert conds_block is not None
    flowables_list = getattr(conds_block, "_content", getattr(conds_block, "flowables", []))
    conds_text = str([el._cellvalues for el in flowables_list if hasattr(el, "_cellvalues")])
    assert "Delivery Address:" in conds_text
    assert "Special export packaging" in conds_text
    # "Ek Alan:" veya "Additional Notes:" başlığı olmamalıdır
    assert "Ek Alan:" not in conds_text
    assert "Additional Notes:" not in conds_text

    # 5. Signatures flowables testi
    sig_table = service._build_signatures(offer, 500)
    sig_text = str(sig_table._cellvalues)
    assert "ACME INTERNATIONAL" in sig_text

    # 6. PDF Dosyası Üretimi
    out = service.generate_pdf(offer, pdf_path)
    assert out.exists()

    doc = QPdfDocument()
    doc.load(str(out))
    assert doc.pageCount() == 1


def test_excel_generation_english_and_additional_notes(tmp_path: Path):
    xlsx_path = tmp_path / "test_en.xlsx"
    service = ExcelService()

    offer = Offer(
        offer_number="20261021-003",
        offer_date=date(2026, 10, 21),
        company_id="endemiks_gida",
        recipient_name="London Food Traders",
        currency_code=CurrencyCode.EUR,
        quantity_unit=QuantityUnit.LITRE,
        vat_enabled=True,
        vat_rate=Decimal("20.00"),
        discount_enabled=True,
        document_language=DocumentLanguage.EN,
        selected_columns=[
            ColumnKey.DESCRIPTION,
            ColumnKey.QUANTITY,
            ColumnKey.UNIT_PRICE,
            ColumnKey.DISCOUNT_TYPE,
            ColumnKey.DISCOUNT_AMOUNT,
            ColumnKey.VAT,
            ColumnKey.TOTAL_AMOUNT,
            ColumnKey.TOTAL_DISCOUNTED,
            ColumnKey.TOTAL_WITH_VAT,
        ],
        items=[
            OfferItem(
                product_code="EU-01",
                description="Cold Pressed Olive Oil",
                quantity=Decimal("50"),
                unit_price=Decimal("40.00"),
                discount_type=DiscountType.AMOUNT,
                discount_value=Decimal("100.00"),
            )
        ],
        sections=[
            OfferSection(
                section_key=SectionKey.PAYMENT_TERMS,
                text="30 days deferred bank transfer.",
                enabled=True,
            ),
            OfferSection(
                section_key=SectionKey.ADDITIONAL_NOTES,
                text="<b>Temperature controlled container shipment.</b>",
                enabled=True,
                rich_text=True,
            ),
        ],
    )

    out = service.generate_excel(offer, xlsx_path)
    assert out.exists()

    wb = openpyxl.load_workbook(str(out), data_only=True)
    assert wb.active.title == "PRICE QUOTATION"
    ws = wb.active

    # Tüm hücre değerlerini toplayalım
    cell_values: list[str] = []
    for row in ws.iter_rows(values_only=True):
        for cell in row:
            if cell is not None:
                cell_values.append(str(cell))

    # İngilizce başlıklar kontrolü
    assert "OFFER NO:" in cell_values
    assert "DATE:" in cell_values
    assert "CURRENCY:" in cell_values
    assert "Customer / Company:" in cell_values
    assert "Description" in cell_values
    assert "Quantity (Liter)" in cell_values
    assert "Amount" in cell_values
    assert "Subtotal:" in cell_values
    assert "Total Discount:" in cell_values
    assert "Grand Total (Incl. VAT):" in cell_values
    assert "OFFER TERMS AND CONDITIONS" in cell_values
    assert "Payment Terms:" in cell_values

    # Ek Alan kontrolü: Sabit başlık ("Ek Alan:" / "Additional Notes:") OLMAMALI, metin ise olmalı
    assert "Temperature controlled container shipment." in cell_values
    assert "Ek Alan:" not in cell_values
    assert "Additional Notes:" not in cell_values


def test_ui_header_widget_and_main_window_language(qtbot):
    from teklif_sablonu.ui.header_widget import HeaderWidget
    from teklif_sablonu.ui.main_window import MainWindow

    hw = HeaderWidget()
    qtbot.addWidget(hw)

    assert hw.get_document_language() == DocumentLanguage.TR
    hw.set_document_language(DocumentLanguage.EN)
    assert hw.get_document_language() == DocumentLanguage.EN
    hw.set_document_language("tr")
    assert hw.get_document_language() == DocumentLanguage.TR

    # MainWindow testi
    mw = MainWindow()
    qtbot.addWidget(mw)

    mw.header_widget.set_document_language(DocumentLanguage.EN)
    offer = mw.to_offer()
    assert offer.document_language == DocumentLanguage.EN

    # Sıfırlama
    mw.clear_offer()
    assert mw.header_widget.get_document_language() == DocumentLanguage.TR
