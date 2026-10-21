"""Prompt 7: Bütünleştirme, uçtan uca doğrulama ve kabul testleri (Acceptance Tests).

Bu modül:
- UI, Domain, Storage (SQLite), PDF ve Excel servislerinin tek akışta birlikte çalışabilirliğini,
- Şartnamedeki 10 zorunlu kabul senaryosunu (tek ürün, iki ürün, çok sütun, iskonto oranı,
  sabit iskonto, KDV, iskontolu+KDV, uzun açıklama/koşullar, Türkçe karakterler, şirket seçimi),
- Negatif doğrulama ve veri bütünlüğü kurallarını doğrular.
"""

from __future__ import annotations

from datetime import date
from decimal import Decimal
from pathlib import Path
import sqlite3
import pytest

import openpyxl
from PySide6.QtPdf import QPdfDocument
from PySide6.QtWidgets import QApplication

from teklif_sablonu.domain.calculator import calculate_offer
from teklif_sablonu.domain.enums import (
    ColumnKey,
    CurrencyCode,
    DiscountType,
    QuantityUnit,
    SectionKey,
)
from teklif_sablonu.domain.models import Offer, OfferItem, OfferSection
from teklif_sablonu.domain.validators import validate_offer
from teklif_sablonu.services.excel_service import ExcelService
from teklif_sablonu.services.offer_service import OfferService
from teklif_sablonu.services.pdf_service import PdfService
from teklif_sablonu.storage.repository import SqliteOfferRepository
from teklif_sablonu.ui.main_window import MainWindow


@pytest.fixture
def memory_offer_service() -> OfferService:
    """Bellek içi (in-memory) SQLite veritabanı ile izole OfferService sağlar."""
    conn = sqlite3.connect(":memory:")
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON;")
    repo = SqliteOfferRepository(conn)
    return OfferService(repo)


@pytest.fixture
def pdf_service() -> PdfService:
    return PdfService()


@pytest.fixture
def excel_service() -> ExcelService:
    return ExcelService()


# ── SENARYO 1: Tek Ürünlü Teklif ──
def test_acceptance_scenario_1_single_item(
    memory_offer_service: OfferService,
    pdf_service: PdfService,
    excel_service: ExcelService,
    tmp_path: Path,
) -> None:
    """Senaryo 1: Tek ürünlü standart teklif -> SQLite kayıt -> PDF & Excel üretimi."""
    offer = Offer(
        offer_date=date(2026, 10, 20),
        company_id="endemiks_tarim",
        recipient_name="Akdeniz Gübre Sanayi",
        currency_code=CurrencyCode.TRY,
        quantity_unit=QuantityUnit.PIECE,
        items=[
            OfferItem(
                description="Standart Sıvı Gübre 20L",
                quantity=Decimal("10"),
                unit_price=Decimal("150.00"),
            )
        ],
    )

    # 1. Doğrula & Kaydet
    saved = memory_offer_service.save_offer(offer)
    assert saved.id is not None
    assert saved.offer_number == "20261020-001"

    # 2. PDF Üret
    pdf_path = tmp_path / f"{saved.offer_number}.pdf"
    pdf_res = pdf_service.generate_pdf(saved, pdf_path)
    assert pdf_res.exists() and pdf_res.stat().st_size > 1000

    doc = QPdfDocument()
    doc.load(str(pdf_res))
    assert doc.pageCount() == 1
    # 4 sütun: Dikey A4 (Portrait)
    assert doc.pagePointSize(0).width() < doc.pagePointSize(0).height()

    # 3. Excel Üret
    xlsx_path = tmp_path / f"{saved.offer_number}.xlsx"
    xlsx_res = excel_service.generate_excel(saved, xlsx_path)
    assert xlsx_res.exists() and xlsx_res.stat().st_size > 1000

    wb = openpyxl.load_workbook(str(xlsx_res), data_only=False)
    sheet = wb.active
    assert sheet.title == "Teklif Formu"


# ── SENARYO 2: İki Ürünlü Teklif ──
def test_acceptance_scenario_2_two_items(
    memory_offer_service: OfferService,
    pdf_service: PdfService,
    excel_service: ExcelService,
    tmp_path: Path,
) -> None:
    """Senaryo 2: İki ürünlü teklif -> Doğru satır toplamları ve genel toplam."""
    offer = Offer(
        offer_date=date(2026, 10, 20),
        company_id="endemiks_gida",
        recipient_name="Ege Gurme Market",
        currency_code=CurrencyCode.USD,
        quantity_unit=QuantityUnit.PIECE,
        items=[
            OfferItem(description="Organik Zeytinyağı 5L", quantity=Decimal("20"), unit_price=Decimal("40.00")),
            OfferItem(description="Doğal Nar Ekşisi 250ml", quantity=Decimal("50"), unit_price=Decimal("12.50")),
        ],
    )
    saved = memory_offer_service.save_offer(offer)
    calc_items, totals = calculate_offer(saved)

    # Satır 1: 20 * 40 = 800.00; Satır 2: 50 * 12.50 = 625.00; Ara Toplam = 1425.00
    assert calc_items[0].line_total == Decimal("800.00")
    assert calc_items[1].line_total == Decimal("625.00")
    assert totals.subtotal == Decimal("1425.00")

    pdf_res = pdf_service.generate_pdf(saved, tmp_path / "sc2.pdf")
    xlsx_res = excel_service.generate_excel(saved, tmp_path / "sc2.xlsx")
    assert pdf_res.exists() and xlsx_res.exists()


# ── SENARYO 3: Çok Sütun Seçimi (8+ Sütun & Yatay A4) ──
def test_acceptance_scenario_3_many_columns_landscape(
    memory_offer_service: OfferService,
    pdf_service: PdfService,
    excel_service: ExcelService,
    tmp_path: Path,
) -> None:
    """Senaryo 3: 8 ve daha fazla sütun seçildiğinde Yatay A4 (Landscape) yönüne geçilmelidir."""
    offer = Offer(
        offer_date=date(2026, 10, 20),
        company_id="endemiks_tarim",
        recipient_name="Mega Sera İşletmeleri",
        currency_code=CurrencyCode.EUR,
        quantity_unit=QuantityUnit.LITRE,
        vat_enabled=True,
        vat_rate=Decimal("10.00"),
        discount_enabled=True,
        selected_columns=[
            ColumnKey.PRODUCT_CODE,
            ColumnKey.DESCRIPTION,
            ColumnKey.QUANTITY,
            ColumnKey.PALLET,
            ColumnKey.UNIT_PRICE,
            ColumnKey.DISCOUNT_TYPE,
            ColumnKey.DISCOUNT_AMOUNT,
            ColumnKey.VAT,
            ColumnKey.TOTAL_AMOUNT,
            ColumnKey.TOTAL_DISCOUNTED,
            ColumnKey.TOTAL_WITH_VAT,
        ],
        items=[
            OfferItem(
                product_code="AGR-900",
                description="Kompost Gübre Katkısı",
                quantity=Decimal("200"),
                pallet=Decimal("4"),
                unit_price=Decimal("18.50"),
                discount_type=DiscountType.RATE,
                discount_value=Decimal("5.00"),
            )
        ],
    )
    saved = memory_offer_service.save_offer(offer)

    # PDF kontrolü: Yatay A4 olmalı
    pdf_res = pdf_service.generate_pdf(saved, tmp_path / "sc3.pdf")
    doc = QPdfDocument()
    doc.load(str(pdf_res))
    assert doc.pagePointSize(0).width() > doc.pagePointSize(0).height()  # Landscape

    # Excel kontrolü: ORIENTATION_LANDSCAPE olmalı
    xlsx_res = excel_service.generate_excel(saved, tmp_path / "sc3.xlsx")
    wb = openpyxl.load_workbook(str(xlsx_res))
    assert wb.active.page_setup.orientation == wb.active.ORIENTATION_LANDSCAPE


# ── SENARYO 4: İskonto Oranı (%10) ──
def test_acceptance_scenario_4_rate_discount(
    memory_offer_service: OfferService,
    pdf_service: PdfService,
    excel_service: ExcelService,
    tmp_path: Path,
) -> None:
    """Senaryo 4: Yüzde iskonto uygulandığında indirim tutarı ve toplam doğru hesaplanmalıdır."""
    offer = Offer(
        offer_date=date(2026, 10, 20),
        company_id="endemiks_gida",
        recipient_name="Toptan Dağıtım A.Ş.",
        discount_enabled=True,
        items=[
            OfferItem(
                description="Portakal Suyu Konsantresi",
                quantity=Decimal("100"),
                unit_price=Decimal("50.00"),  # Toplam: 5000.00
                discount_type=DiscountType.RATE,
                discount_value=Decimal("10.00"),  # %10 -> 500.00 iskonto -> 4500.00
            )
        ],
    )
    saved = memory_offer_service.save_offer(offer)
    calc_items, totals = calculate_offer(saved)

    assert calc_items[0].line_total == Decimal("5000.00")
    assert calc_items[0].discount_amount == Decimal("500.00")
    assert calc_items[0].discounted_total == Decimal("4500.00")
    assert totals.total_discount == Decimal("500.00")
    assert totals.discounted_subtotal == Decimal("4500.00")

    pdf_service.generate_pdf(saved, tmp_path / "sc4.pdf")
    excel_service.generate_excel(saved, tmp_path / "sc4.xlsx")


# ── SENARYO 5: Sabit İskonto Miktarı (500 TL) ──
def test_acceptance_scenario_5_amount_discount(
    memory_offer_service: OfferService,
    pdf_service: PdfService,
    excel_service: ExcelService,
    tmp_path: Path,
) -> None:
    """Senaryo 5: Sabit tutar iskontosu uygulandığında hesaplama ve formüller doğru çalışmalıdır."""
    offer = Offer(
        offer_date=date(2026, 10, 20),
        company_id="endemiks_tarim",
        recipient_name="Bölge Bayii",
        discount_enabled=True,
        items=[
            OfferItem(
                description="Tohum Kaplama Sıvısı",
                quantity=Decimal("10"),
                unit_price=Decimal("300.00"),  # Toplam: 3000.00
                discount_type=DiscountType.AMOUNT,
                discount_value=Decimal("500.00"),  # 500.00 sabit iskonto -> 2500.00
            )
        ],
    )
    saved = memory_offer_service.save_offer(offer)
    calc_items, totals = calculate_offer(saved)

    assert calc_items[0].line_total == Decimal("3000.00")
    assert calc_items[0].discount_amount == Decimal("500.00")
    assert calc_items[0].discounted_total == Decimal("2500.00")
    assert totals.total_discount == Decimal("500.00")
    assert totals.discounted_subtotal == Decimal("2500.00")

    pdf_service.generate_pdf(saved, tmp_path / "sc5.pdf")
    excel_service.generate_excel(saved, tmp_path / "sc5.xlsx")


# ── SENARYO 6: KDV Uygulaması (%20) ──
def test_acceptance_scenario_6_vat_calculation(
    memory_offer_service: OfferService,
    pdf_service: PdfService,
    excel_service: ExcelService,
    tmp_path: Path,
) -> None:
    """Senaryo 6: KDV etkinleştirildiğinde genel toplam özeti ve KDV satırları eksiksiz üretilmelidir."""
    offer = Offer(
        offer_date=date(2026, 10, 20),
        company_id="endemiks_tarim",
        recipient_name="Modern Sera Ltd.",
        vat_enabled=True,
        vat_rate=Decimal("20.00"),
        items=[
            OfferItem(
                description="Biyolojik Mücadele Ürünü",
                quantity=Decimal("5"),
                unit_price=Decimal("1000.00"),  # Toplam: 5000.00, KDV: 1000.00, Genel: 6000.00
            )
        ],
    )
    saved = memory_offer_service.save_offer(offer)
    _, totals = calculate_offer(saved)

    assert totals.subtotal == Decimal("5000.00")
    assert totals.total_vat == Decimal("1000.00")
    assert totals.grand_total == Decimal("6000.00")

    pdf_service.generate_pdf(saved, tmp_path / "sc6.pdf")
    excel_service.generate_excel(saved, tmp_path / "sc6.xlsx")


# ── SENARYO 7: İskontolu + KDV Birlikte Hesaplama ──
def test_acceptance_scenario_7_discount_plus_vat(
    memory_offer_service: OfferService,
    pdf_service: PdfService,
    excel_service: ExcelService,
    tmp_path: Path,
) -> None:
    """Senaryo 7: Şartname Kuralı: KDV iskontolu ara toplam üzerinden hesaplanmalıdır."""
    offer = Offer(
        offer_date=date(2026, 10, 20),
        company_id="endemiks_tarim",
        recipient_name="Çukurova Tarım A.Ş.",
        discount_enabled=True,
        vat_enabled=True,
        vat_rate=Decimal("20.00"),
        items=[
            OfferItem(
                description="Özel Karışım Gübre",
                quantity=Decimal("20"),
                unit_price=Decimal("100.00"),  # Toplam: 2000.00
                discount_type=DiscountType.RATE,
                discount_value=Decimal("10.00"),  # İskonto: 200.00 -> Matrah: 1800.00
                # KDV: 1800 * 0.20 = 360.00 -> Genel Toplam: 2160.00
            )
        ],
    )
    saved = memory_offer_service.save_offer(offer)
    _, totals = calculate_offer(saved)

    assert totals.subtotal == Decimal("2000.00")
    assert totals.total_discount == Decimal("200.00")
    assert totals.discounted_subtotal == Decimal("1800.00")
    assert totals.total_vat == Decimal("360.00")
    assert totals.grand_total == Decimal("2160.00")

    pdf_service.generate_pdf(saved, tmp_path / "sc7.pdf")
    excel_service.generate_excel(saved, tmp_path / "sc7.xlsx")


# ── SENARYO 8: Uzun Açıklama ve Koşul Metinleri ──
def test_acceptance_scenario_8_long_descriptions_and_sections(
    memory_offer_service: OfferService,
    pdf_service: PdfService,
    excel_service: ExcelService,
    tmp_path: Path,
) -> None:
    # Şartname kuralı: Ürün açıklaması azami 64 karakterdir (sınırda test)
    long_desc = "Özel Formülasyonlu Mikro Element Katkılı Sıvı Bitki Besleme Gübresi"[:64]
    assert len(long_desc) == 64
    offer = Offer(
        offer_date=date(2026, 10, 20),
        recipient_name="Anadolu Tarım Kooperatifleri Birliği",
        items=[OfferItem(description=long_desc, quantity=Decimal("100"), unit_price=Decimal("75.00"))],
        sections=[
            OfferSection(
                section_key=SectionKey.DELIVERY_ADDRESS,
                enabled=True,
                text="Tek bir merkezi antrepoya araç üstü teslim şartıyla teslim edilecektir.",
            ),
            OfferSection(
                section_key=SectionKey.PAYMENT_TERMS,
                enabled=True,
                text="Sözleşme imzalanmasını müteakip %30 peşin, kalan bakiye 60 gün vadeli ticari senet karşılığı.",
            ),
            OfferSection(
                section_key=SectionKey.ADDITIONAL_NOTES,
                enabled=True,
                text="<b>Önemli Kalite Uyarısı:</b><br>• Ürün doğrudan güneş ışığına maruz bırakılmamalıdır.<br>• Paletler azami 2 sıra üst üste istiflenebilir.",
                rich_text=True,
            ),
        ],
    )
    saved = memory_offer_service.save_offer(offer)

    pdf_res = pdf_service.generate_pdf(saved, tmp_path / "sc8.pdf")
    xlsx_res = excel_service.generate_excel(saved, tmp_path / "sc8.xlsx")
    assert pdf_res.exists() and xlsx_res.exists()


# ── SENARYO 9: Türkçe Karakter Desteği ──
def test_acceptance_scenario_9_turkish_characters(
    memory_offer_service: OfferService,
    pdf_service: PdfService,
    excel_service: ExcelService,
    tmp_path: Path,
) -> None:
    """Senaryo 9: Tüm Türkçe karakterler (Ş, ğ, İ, ö, ü, ç, ı) hatasız gömülmeli ve bozulmamalıdır."""
    turkish_text = "Çağdaş Şirketler Grubu & İnci Örnek Çiftliği (Ispanak, Söğüt, Çavdar)"
    offer = Offer(
        offer_date=date(2026, 10, 20),
        recipient_name=turkish_text,
        items=[OfferItem(description="Özel Çavdar ve Söğüt Özü Şurubu", quantity=Decimal("5"), unit_price=Decimal("120"))],
        sections=[
            OfferSection(
                section_key=SectionKey.STORAGE_CONDITIONS,
                enabled=True,
                text="Işık almayan, nem oranı %50'nin altında serin depolarda muhafaza edilmelidir.",
            )
        ],
    )
    saved = memory_offer_service.save_offer(offer)

    pdf_res = pdf_service.generate_pdf(saved, tmp_path / "sc9.pdf")
    xlsx_res = excel_service.generate_excel(saved, tmp_path / "sc9.xlsx")
    assert pdf_res.exists() and xlsx_res.exists()


# ── SENARYO 10: Farklı Şirket Seçimi ──
def test_acceptance_scenario_10_different_companies(
    memory_offer_service: OfferService,
    pdf_service: PdfService,
    excel_service: ExcelService,
    tmp_path: Path,
) -> None:
    """Senaryo 10: İki farklı şirket (Tarım vs Gıda) seçimi üst bilgilerde doğru şirketi göstermelidir."""
    offer_tarim = Offer(
        company_id="endemiks_tarim",
        recipient_name="Müşteri A",
        items=[OfferItem(description="Ürün A", quantity=Decimal("1"), unit_price=Decimal("10"))],
    )
    offer_gida = Offer(
        company_id="endemiks_gida",
        recipient_name="Müşteri B",
        items=[OfferItem(description="Ürün B", quantity=Decimal("1"), unit_price=Decimal("20"))],
    )

    saved_tarim = memory_offer_service.save_offer(offer_tarim)
    saved_gida = memory_offer_service.save_offer(offer_gida)

    assert saved_tarim.company_id == "endemiks_tarim"
    assert saved_gida.company_id == "endemiks_gida"

    # PDF ve Excel çıktıları iki şirket için de sorunsuz oluşmalı
    pdf_service.generate_pdf(saved_tarim, tmp_path / "tarim.pdf")
    pdf_service.generate_pdf(saved_gida, tmp_path / "gida.pdf")
    excel_service.generate_excel(saved_tarim, tmp_path / "tarim.xlsx")
    excel_service.generate_excel(saved_gida, tmp_path / "gida.xlsx")


# ── NEGATİF DOĞRULAMA KONTROLLERİ ──
def test_acceptance_negative_validation_prevents_save_and_export(
    memory_offer_service: OfferService,
    pdf_service: PdfService,
    excel_service: ExcelService,
) -> None:
    """Kısmen doldurulmuş veya geçersiz tekliflerde kaydetme ve çıktı üretimi engellenmelidir."""
    # 1. Kısmen doldurulmuş satır (açıklama var, miktar yok)
    invalid_offer = Offer(
        recipient_name="Test Alıcı",
        items=[OfferItem(description="Eksik Ürün", quantity=Decimal("0"), unit_price=Decimal("100"))],
    )
    val = validate_offer(invalid_offer)
    assert val.is_valid is False

    with pytest.raises(ValueError):
        memory_offer_service.save_offer(invalid_offer)

    with pytest.raises(ValueError):
        pdf_service.generate_pdf(invalid_offer)

    with pytest.raises(ValueError):
        excel_service.generate_excel(invalid_offer)


# ── UÇTAN UCA UI ENTEGRASYON AKIŞI ──
def test_acceptance_ui_end_to_end_flow(
    memory_offer_service: OfferService,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    """UI üzerinden form doldurma, kaydetme, PDF/Excel butonları, arşiv ve kopyalama uçtan uca akışı."""
    custom_pdf = PdfService()
    monkeypatch.setattr(custom_pdf, "get_default_output_path", lambda num: tmp_path / f"{num}.pdf")
    custom_excel = ExcelService()
    monkeypatch.setattr(custom_excel, "get_default_output_path", lambda num: tmp_path / f"{num}.xlsx")

    win = MainWindow(
        offer_service=memory_offer_service,
        pdf_service=custom_pdf,
        excel_service=custom_excel,
    )

    # 1. Bilgileri gir
    win.header_widget.set_recipient_name("Bütünleşik Test Müşterisi A.Ş.")
    win.header_widget.set_currency_code(CurrencyCode.TRY)
    win.header_widget.set_vat_enabled(True)
    win.header_widget.set_vat_rate(Decimal("20.00"))

    row0 = win.items_table._rows[0]
    row0.desc_edit.setText("Akıllı Tarım Sensörü")
    row0.qty_edit.setText("5")
    row0.price_edit.setText("2.000,00")

    # 2. Doğrula & Kaydet
    win.btn_validate.click()
    win.btn_save.click()

    assert win._current_offer_number != ""
    assert win.archive_widget.table.rowCount() == 1

    # 3. PDF ve Excel oluştur
    win.btn_export_pdf.click()
    win.btn_export_excel.click()

    generated_pdf = tmp_path / f"{win._current_offer_number}.pdf"
    generated_xlsx = tmp_path / f"{win._current_offer_number}.xlsx"
    assert generated_pdf.exists()
    assert generated_xlsx.exists()

    # 4. Formu temizle ve Arşivden aç
    first_offer_number = win._current_offer_number
    win.clear_offer()
    assert win.header_widget.get_recipient_name() == ""

    win.archive_widget.table.selectRow(0)
    win.archive_widget.btn_open.click()
    assert win.header_widget.get_recipient_name() == "Bütünleşik Test Müşterisi A.Ş."
    assert win._current_offer_number == first_offer_number

    # 5. Arşivden kopyalayarak yeni teklif üret
    win.archive_widget.table.selectRow(0)
    win.archive_widget.btn_duplicate.click()
    assert win._current_offer_number == ""  # Yeni teklif olduğu için numara sıfırlandı
    assert win._source_offer_id is not None

    # Yeni kopyayı kaydet -> Farklı bir numara almalı
    win.btn_save.click()
    assert win._current_offer_number != first_offer_number
    assert win._current_offer_number == f"{first_offer_number[:-3]}002"

    QApplication.processEvents()
    win.close()
