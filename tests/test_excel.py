"""Excel dışa aktarım servisi birim ve entegrasyon testleri."""

from __future__ import annotations

from datetime import date
from decimal import Decimal
from pathlib import Path
import pytest

import openpyxl

from teklif_sablonu.domain.enums import (
    ColumnKey,
    CurrencyCode,
    DiscountType,
    QuantityUnit,
    SectionKey,
)
from teklif_sablonu.domain.models import Offer, OfferItem, OfferSection
from teklif_sablonu.services.excel_service import ExcelService, strip_html_tags


@pytest.fixture
def excel_service() -> ExcelService:
    return ExcelService()


def test_excel_generation_standard(excel_service: ExcelService, tmp_path: Path) -> None:
    """Temel teklif verisi geçerli bir .xlsx dosyası olarak üretilmelidir."""
    out_xlsx = tmp_path / "test_standard.xlsx"

    offer = Offer(
        offer_number="20261020-001",
        offer_date=date(2026, 10, 20),
        company_id="endemiks_tarim",
        recipient_name="Antalya Tarım A.Ş.",
        currency_code=CurrencyCode.TRY,
        quantity_unit=QuantityUnit.PIECE,
        vat_enabled=False,
        discount_enabled=False,
        brand_logo_visible=True,
        selected_columns=[
            ColumnKey.DESCRIPTION,
            ColumnKey.QUANTITY,
            ColumnKey.UNIT_PRICE,
            ColumnKey.TOTAL_AMOUNT,
        ],
        items=[
            OfferItem(
                description="Organik Sıvı Gübre 1L",
                quantity=Decimal("50"),
                unit_price=Decimal("120.00"),
            )
        ],
    )

    result_path = excel_service.generate_excel(offer, out_xlsx)
    assert result_path.exists()
    assert result_path.stat().st_size > 1000

    # openpyxl ile oku ve doğrula
    wb = openpyxl.load_workbook(str(result_path), data_only=False)
    sheet = wb.active
    assert sheet.title == "Teklif Formu"

    # Şirket ve alıcı bilgisi kontrolü
    found_recipient = False
    for r in range(1, 15):
        for c in range(1, 6):
            val = sheet.cell(row=r, column=c).value
            if val and "Antalya Tarım A.Ş." in str(val):
                found_recipient = True
                break
    assert found_recipient is True


def test_excel_formulas_and_numeric_types(excel_service: ExcelService, tmp_path: Path) -> None:
    """Excel çıktısında tutarlar formüllü, miktar ve birim fiyat sayısal olmalıdır."""
    out_xlsx = tmp_path / "test_formulas.xlsx"

    offer = Offer(
        offer_number="20261020-002",
        offer_date=date(2026, 10, 20),
        company_id="endemiks_gida",
        recipient_name="İstanbul Restoran Ltd.",
        currency_code=CurrencyCode.USD,
        quantity_unit=QuantityUnit.PIECE,
        vat_enabled=False,
        discount_enabled=False,
        items=[
            OfferItem(
                description="Zeytinyağı 5L",
                quantity=Decimal("10"),
                unit_price=Decimal("25.50"),
            ),
            OfferItem(
                description="Nar Ekşisi 250ml",
                quantity=Decimal("20"),
                unit_price=Decimal("15.00"),
            ),
        ],
    )

    result_path = excel_service.generate_excel(offer, out_xlsx)
    wb = openpyxl.load_workbook(str(result_path), data_only=False)
    sheet = wb.active

    # Tablo başlık satırını bul
    header_row = -1
    for r in range(1, 20):
        if sheet.cell(row=r, column=2).value in ("Ürün Açıklaması", "Ürün Kodu"):
            header_row = r
            break
    assert header_row != -1

    # İlk veri satırı
    data_r1 = header_row + 1
    data_r2 = header_row + 2

    # Miktar ve fiyat sayısal olmalı (string değil)
    # Varsayılan görünür sütunlar: Açıklama(B), Miktar(C), Birim Fiyat(D), Toplam(E)
    qty_val = sheet.cell(row=data_r1, column=3).value
    price_val = sheet.cell(row=data_r1, column=4).value
    total_formula = sheet.cell(row=data_r1, column=5).value

    assert isinstance(qty_val, (int, float))
    assert qty_val == 10.0
    assert isinstance(price_val, (int, float))
    assert price_val == 25.50

    # Toplam hücre formülü =C...*D... olmalı
    assert isinstance(total_formula, str)
    assert total_formula.startswith("=")
    assert "*" in total_formula

    # Ara toplam hücresinde SUM formülü olmalı
    found_sum = False
    for r in range(data_r2 + 1, sheet.max_row + 1):
        for c in range(1, 10):
            val = sheet.cell(row=r, column=c).value
            if isinstance(val, str) and "=SUM(" in val:
                found_sum = True
                break
    assert found_sum is True


def test_excel_discount_and_vat_features(excel_service: ExcelService, tmp_path: Path) -> None:
    """İskonto ve KDV aktif olduğunda formüller ve toplam kartı doğru oluşmalıdır."""
    out_xlsx = tmp_path / "test_disc_vat.xlsx"

    offer = Offer(
        offer_number="20261020-003",
        offer_date=date(2026, 10, 20),
        company_id="endemiks_tarim",
        recipient_name="Çiftçi Kooperatifi",
        currency_code=CurrencyCode.EUR,
        quantity_unit=QuantityUnit.LITRE,
        vat_enabled=True,
        vat_rate=Decimal("20.00"),
        discount_enabled=True,
        selected_columns=[
            ColumnKey.DESCRIPTION,
            ColumnKey.QUANTITY,
            ColumnKey.UNIT_PRICE,
            ColumnKey.DISCOUNT_TYPE,
            ColumnKey.DISCOUNT_AMOUNT,
            ColumnKey.VAT,
            ColumnKey.TOTAL_AMOUNT,
            ColumnKey.TOTAL_DISCOUNTED,
            ColumnKey.TOTAL_WITH_VAT,
        ],
        items=[
            OfferItem(
                description="Damla Sulama Gübresi",
                quantity=Decimal("100"),
                unit_price=Decimal("40.00"),
                discount_type=DiscountType.RATE,
                discount_value=Decimal("10.00"),  # %10 iskonto
            )
        ],
    )

    result_path = excel_service.generate_excel(offer, out_xlsx)
    wb = openpyxl.load_workbook(str(result_path), data_only=False)
    sheet = wb.active

    # Sayfa yönü çok sütunlu olduğundan yatay (Landscape) olmalı
    assert sheet.page_setup.orientation == sheet.ORIENTATION_LANDSCAPE

    # KDV ve İskonto metinlerinin tabloda yer aldığını doğrula
    cell_values = [str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value is not None]
    assert any("Toplam İskonto:" in v for v in cell_values)
    assert any("Toplam KDV (%20):" in v for v in cell_values)
    assert any("KDV Dâhil Genel Toplam:" in v for v in cell_values)


def test_excel_filters_empty_columns(excel_service: ExcelService) -> None:
    """Seçili olsa dahi tüm satırlarda tamamen boş kalan sütunlar Excel'den elenmelidir."""
    offer = Offer(
        selected_columns=[
            ColumnKey.PRODUCT_CODE,  # Boş kalacak
            ColumnKey.DESCRIPTION,
            ColumnKey.QUANTITY,
            ColumnKey.PALLET,        # Boş kalacak
            ColumnKey.UNIT_PRICE,
            ColumnKey.TOTAL_AMOUNT,
        ],
        items=[
            OfferItem(
                product_code="",
                description="Test Ürün",
                quantity=Decimal("5"),
                pallet=None,
                unit_price=Decimal("10"),
            )
        ],
    )

    visible = excel_service.filter_visible_columns(offer, offer.items)
    assert ColumnKey.PRODUCT_CODE not in visible
    assert ColumnKey.PALLET not in visible
    assert ColumnKey.DESCRIPTION in visible
    assert ColumnKey.QUANTITY in visible
    assert ColumnKey.UNIT_PRICE in visible
    assert ColumnKey.TOTAL_AMOUNT in visible


def test_excel_sections_and_html_stripping(excel_service: ExcelService, tmp_path: Path) -> None:
    """Ticari koşullar ve Ek Alan zengin metinleri HTML etiketlerinden arındırılarak eklenmelidir."""
    out_xlsx = tmp_path / "test_sections.xlsx"

    offer = Offer(
        offer_number="20261020-004",
        offer_date=date(2026, 10, 20),
        recipient_name="Mega Market A.Ş.",
        items=[OfferItem(description="Domates Salçası", quantity=Decimal("20"), unit_price=Decimal("30"))],
        sections=[
            OfferSection(
                section_key=SectionKey.PAYMENT_TERMS,
                enabled=True,
                text="Vade: 30 gün",
            ),
            OfferSection(
                section_key=SectionKey.ADDITIONAL_NOTES,
                enabled=True,
                text="<b>Önemli:</b> Paletli sevkiyat.<br>Serin yerde tutunuz.",
                rich_text=True,
            ),
            OfferSection(
                section_key=SectionKey.DELIVERY_METHOD,
                enabled=False,  # Pasif koşul Excel'de yer almamalı
                text="Fabrika Teslim",
            ),
        ],
    )

    result_path = excel_service.generate_excel(offer, out_xlsx)
    wb = openpyxl.load_workbook(str(result_path), data_only=False)
    sheet = wb.active

    all_texts = [str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value is not None]
    # Etkin koşullar mevcut olmalı
    assert any("Ödeme Şartları:" in t for t in all_texts)
    assert any("Vade: 30 gün" in t for t in all_texts)
    # HTML etiketleri arındırılmış olmalı
    assert not any("<b>" in t for t in all_texts)
    assert not any("<br>" in t for t in all_texts)
    assert any("Paletli sevkiyat." in t for t in all_texts)
    # Pasif koşul yer almamalı
    assert not any("Fabrika Teslim" in t for t in all_texts)


def test_excel_validation_failure_raises(excel_service: ExcelService) -> None:
    """Geçersiz tekliflerde Excel üretimi hata fırlatarak durmalıdır."""
    invalid_offer = Offer(
        recipient_name="",  # Zorunlu alan eksik
        items=[],           # Ürün yok
    )
    with pytest.raises(ValueError):
        excel_service.generate_excel(invalid_offer)


def test_strip_html_tags_utility() -> None:
    """strip_html_tags fonksiyonu etiketleri doğru dönüştürmelidir."""
    assert strip_html_tags("") == ""
    assert strip_html_tags("<p>Paragraf 1</p><p>Paragraf 2</p>") == "Paragraf 1\nParagraf 2"
    assert strip_html_tags("Satır 1<br/>Satır 2") == "Satır 1\nSatır 2"
    assert strip_html_tags("<b>Kalın</b> ve <i>İtalik</i>") == "Kalın ve İtalik"


def test_excel_without_total_amount_unit_price_mode(excel_service: ExcelService, tmp_path: Path) -> None:
    """Toplam Tutar kapatıldığında Excel çıktısında ne Toplam Tutar sütunu ne de özet kartı olmalıdır."""
    out_xlsx = tmp_path / "test_no_total.xlsx"
    offer = Offer(
        offer_number="20261029-E01",
        offer_date=date(2026, 10, 29),
        recipient_name="Birim Fiyat Firması",
        currency_code=CurrencyCode.TRY,
        selected_columns=[
            ColumnKey.DESCRIPTION,
            ColumnKey.QUANTITY,
            ColumnKey.UNIT_PRICE,
        ],
        items=[
            OfferItem(description="Kalem 1", quantity=Decimal("1"), unit_price=Decimal("10.00")),
            OfferItem(description="Kalem 2", quantity=Decimal("1"), unit_price=Decimal("20.00")),
        ],
    )
    res = excel_service.generate_excel(offer, out_xlsx)
    assert res.exists()

    wb = openpyxl.load_workbook(res)
    ws = wb.active
    all_values = [cell.value for row in ws.iter_rows() for cell in row if cell.value is not None]
    assert "Ara Toplam:" not in all_values
    assert "Genel Toplam:" not in all_values
    assert "Toplam Tutar" not in [str(v) for v in all_values]
