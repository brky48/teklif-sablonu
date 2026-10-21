"""Excel dışa aktarım servisi — openpyxl ile düzenlenebilir .xlsx üretimi."""

from __future__ import annotations

import os
from pathlib import Path
import re
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from teklif_sablonu.config import ASSETS_DIR, BRANDS_DIR, CompanyInfo, load_companies
from teklif_sablonu.domain.calculator import calculate_offer
from teklif_sablonu.domain.enums import (
    FIXED_COLUMN_ORDER,
    ColumnKey,
    CurrencyCode,
    DiscountType,
    DocumentLanguage,
    PageOrientation,
    QuantityUnit,
    SectionKey,
)
from teklif_sablonu.domain.models import Offer, OfferItem
from teklif_sablonu.domain.validators import validate_offer


def get_desktop_dir() -> Path:
    """Kullanıcının gerçek Masaüstü (Desktop) dizinini bulur ve döndürür."""
    desktop = Path.home() / "Desktop"
    if desktop.exists() and os.access(desktop, os.W_OK):
        return desktop
    return Path.home()


from teklif_sablonu.services.html_utils import qt_html_to_plain_text


def strip_html_tags(text: str) -> str:
    """Ek alan veya zengin metindeki HTML etiketlerini temizler."""
    return qt_html_to_plain_text(text)


class ExcelService:
    """Teklif verilerinden düzenlenebilir, formüllü Excel (.xlsx) oluşturan servis."""

    # Kurumsal renk paleti (ReportLab temasıyla uyumlu)
    PRIMARY_COLOR = "1E3A8A"      # Koyu Lacivert (Başlık arka planı)
    SECONDARY_COLOR = "0F5132"    # Alternatif Kurumsal Koyu Yeşil
    ACCENT_BG = "F3F4F6"          # Açık Gri (Toplamlar ve şeritler)
    HEADER_TEXT_COLOR = "FFFFFF"  # Beyaz
    BORDER_COLOR = "D1D5DB"       # İnce Gri Çerçeve

    def __init__(self) -> None:
        pass

    def _get_company(self, company_id: str) -> CompanyInfo:
        companies = load_companies()
        for c in companies:
            if c.id == company_id:
                return c
        return companies[0]

    def get_default_output_path(self, offer_number: str) -> Path:
        """Varsayılan kayıt yolunu (<Masaüstü>/<TeklifNo>.xlsx) döndürür."""
        clean_num = offer_number or "Teklif"
        return get_desktop_dir() / f"{clean_num}.xlsx"

    def filter_visible_columns(self, offer: Offer, items: list[OfferItem]) -> list[ColumnKey]:
        """Şartname Bölüm 6.2 kuralına göre gösterilecek sütunları filtreler:

        - Seçilmeyen veya tüm ürün satırlarında tamamen boş kalan sütunlar çıkarılır.
        - Zorunlu sütunlar (Açıklama, Miktar, Fiyat, Toplam) daima korunur.
        - Sabit sıra (FIXED_COLUMN_ORDER) kesinlikle korunur.
        """
        visible: list[ColumnKey] = []

        for col in FIXED_COLUMN_ORDER:
            # Kullanıcı tarafından seçilmemişse atla
            if col not in offer.selected_columns:
                continue

            # Zorunlu sütunlar daima kalır
            if col.is_mandatory:
                visible.append(col)
                continue

            # Diğer opsiyonel sütunlar: En az bir satırda değer var mı kontrol et
            has_value = False
            for it in items:
                if col == ColumnKey.PRODUCT_CODE and it.product_code and it.product_code.strip():
                    has_value = True
                    break
                elif col == ColumnKey.PALLET and it.pallet is not None and it.pallet > 0:
                    has_value = True
                    break
                elif col == ColumnKey.DISCOUNT_TYPE and offer.discount_enabled and it.discount_type != DiscountType.NONE:
                    has_value = True
                    break
                elif col == ColumnKey.DISCOUNT_AMOUNT and offer.discount_enabled and it.discount_value > 0:
                    has_value = True
                    break
                elif col == ColumnKey.VAT and offer.vat_enabled and offer.vat_rate > 0:
                    has_value = True
                    break
                elif col == ColumnKey.TOTAL_DISCOUNTED and offer.discount_enabled:
                    has_value = True
                    break
                elif col == ColumnKey.TOTAL_WITH_VAT and offer.vat_enabled:
                    has_value = True
                    break

            if has_value:
                visible.append(col)

        return visible

    def generate_excel(self, offer: Offer, output_path: Path | None = None) -> Path:
        """Domain Offer modelinden düzenlenebilir, formüllü Excel dosyası oluşturur.

        Args:
            offer: Kaydedilmiş veya doğrulanmış teklif modeli.
            output_path: Çıktı dosya yolu. Belirtilmezse masaüstü yolu kullanılır.

        Returns:
            Oluşturulan Excel dosyasının Path nesnesi.
        """
        # 1. Domain seviyesinde kuralları doğrula
        validation = validate_offer(offer)
        validation.raise_if_invalid()

        target_path = output_path or self.get_default_output_path(offer.offer_number)
        target_path.parent.mkdir(parents=True, exist_ok=True)

        valid_items = [it for it in offer.items if not it.is_empty()]
        _, offer_totals = calculate_offer(offer)
        visible_cols = self.filter_visible_columns(offer, valid_items)
        is_en = offer.document_language == DocumentLanguage.EN

        # 2. Workbook ve Worksheet başlat
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "PRICE QUOTATION" if is_en else "Teklif Formu"

        # Sayfa Yönü: Kullanıcı tercihi veya sütun sayısına göre (7 veya daha az sütun Dikey, 8+ sütun Yatay)
        is_landscape = False
        if offer.page_orientation == PageOrientation.LANDSCAPE or offer.page_orientation == "landscape":
            is_landscape = True
        elif offer.page_orientation == PageOrientation.PORTRAIT or offer.page_orientation == "portrait":
            is_landscape = False
        else:
            is_landscape = len(visible_cols) > 7

        ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE if is_landscape else ws.ORIENTATION_PORTRAIT
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0

        # Stil Yardımcıları
        thin_border = Side(style="thin", color=self.BORDER_COLOR)
        cell_border = Border(left=thin_border, right=thin_border, top=thin_border, bottom=thin_border)
        header_fill = PatternFill(start_color=self.PRIMARY_COLOR, end_color=self.PRIMARY_COLOR, fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color=self.HEADER_TEXT_COLOR)
        bold_font = Font(name="Calibri", size=10, bold=True)
        regular_font = Font(name="Calibri", size=10)
        title_font = Font(name="Calibri", size=14, bold=True, color=self.PRIMARY_COLOR)
        muted_font = Font(name="Calibri", size=9, color="555555")

        company = self._get_company(offer.company_id)

        # 3. Şirket ve Teklif Üst Bilgileri (Satır 2 - 7)
        # Sol Taraf: Şirket Bilgileri
        ws.cell(row=2, column=2, value=company.display_name).font = title_font
        current_meta_row = 3
        if company.tax_office and company.tax_id:
            ws.cell(row=current_meta_row, column=2, value=f"{company.tax_office} V.D. - VKN: {company.tax_id}").font = muted_font
            current_meta_row += 1
        for addr in company.address_lines:
            ws.cell(row=current_meta_row, column=2, value=addr).font = muted_font
            current_meta_row += 1

        # Sağ Taraf: Teklif Numarası ve Tarihi (Tablonun son sütununa hizalı)
        right_col = max(len(visible_cols) + 1, 6)
        ws.cell(row=2, column=right_col - 1, value="OFFER NO:" if is_en else "TEKLİF NO:").font = bold_font
        ws.cell(row=2, column=right_col - 1).alignment = Alignment(horizontal="right")
        ws.cell(row=2, column=right_col, value=offer.offer_number or "—").font = bold_font

        ws.cell(row=3, column=right_col - 1, value="DATE:" if is_en else "TARİH:").font = bold_font
        ws.cell(row=3, column=right_col - 1).alignment = Alignment(horizontal="right")
        date_str = offer.offer_date.strftime("%d.%m.%Y") if offer.offer_date else "—"
        ws.cell(row=3, column=right_col, value=date_str).font = regular_font

        ws.cell(row=4, column=right_col - 1, value="CURRENCY:" if is_en else "PARA BİRİMİ:").font = muted_font
        ws.cell(row=4, column=right_col - 1).alignment = Alignment(horizontal="right")
        ws.cell(row=4, column=right_col, value=offer.currency_code.value).font = muted_font

        # Marka Logosu (Juresh) — Eğer aktifse sağ üst alana ekle
        if offer.brand_logo_visible:
            juresh_logo = BRANDS_DIR / "juresh_logo.png"
            if juresh_logo.exists():
                try:
                    img = openpyxl.drawing.image.Image(str(juresh_logo))
                    img.width = 110
                    img.height = 42
                    ws.add_image(img, f"{get_column_letter(right_col)}5")
                except Exception:
                    pass

        # 4. Müşteri (Alıcı) Bilgileri
        client_row = max(current_meta_row + 1, 7)
        ws.cell(row=client_row, column=2, value="Customer / Company:" if is_en else "Sayın / Firma:").font = bold_font
        ws.cell(row=client_row, column=3, value=offer.recipient_name).font = Font(name="Calibri", size=11, bold=True)

        # 5. Ürün Tablosu Başlık Satırı
        table_header_row = client_row + 2
        col_letter_map: dict[ColumnKey, str] = {}

        for idx, col_key in enumerate(visible_cols, start=2):
            col_letter = get_column_letter(idx)
            col_letter_map[col_key] = col_letter
            cell = ws.cell(row=table_header_row, column=idx)
            cell.value = self._get_column_header(col_key, offer)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = cell_border
        ws.row_dimensions[table_header_row].height = 28

        # 6. Ürün Veri Satırları & Formüller
        start_data_row = table_header_row + 1
        num_format_currency = f"#,##0.00 \"{offer.currency_code.symbol}\""
        num_format_rate = "0.00%"

        qty_letter = col_letter_map.get(ColumnKey.QUANTITY, "")
        price_letter = col_letter_map.get(ColumnKey.UNIT_PRICE, "")
        total_letter = col_letter_map.get(ColumnKey.TOTAL_AMOUNT, "")
        disc_val_letter = col_letter_map.get(ColumnKey.DISCOUNT_AMOUNT, "")
        disc_total_letter = col_letter_map.get(ColumnKey.TOTAL_DISCOUNTED, "")
        vat_letter = col_letter_map.get(ColumnKey.VAT, "")

        for r_offset, item in enumerate(valid_items):
            current_r = start_data_row + r_offset
            ws.row_dimensions[current_r].height = 22

            for col_key in visible_cols:
                c_idx = openpyxl.utils.column_index_from_string(col_letter_map[col_key])
                cell = ws.cell(row=current_r, column=c_idx)
                cell.font = regular_font
                cell.border = cell_border

                if col_key == ColumnKey.PRODUCT_CODE:
                    cell.value = item.product_code or ""
                    cell.alignment = Alignment(horizontal="center", vertical="center")

                elif col_key == ColumnKey.DESCRIPTION:
                    cell.value = item.description
                    cell.alignment = Alignment(horizontal="left", vertical="center")

                elif col_key == ColumnKey.QUANTITY:
                    cell.value = float(item.quantity)
                    cell.number_format = "#,##0.00" if (item.quantity % 1 != 0) else "#,##0"
                    cell.alignment = Alignment(horizontal="right", vertical="center")

                elif col_key == ColumnKey.PALLET:
                    if item.pallet is not None:
                        cell.value = float(item.pallet)
                        cell.number_format = "#,##0.00" if (item.pallet % 1 != 0) else "#,##0"
                    cell.alignment = Alignment(horizontal="right", vertical="center")

                elif col_key == ColumnKey.UNIT_PRICE:
                    cell.value = float(item.unit_price)
                    cell.number_format = num_format_currency
                    cell.alignment = Alignment(horizontal="right", vertical="center")

                elif col_key == ColumnKey.DISCOUNT_TYPE:
                    if offer.discount_enabled:
                        if item.discount_type == DiscountType.RATE:
                            cell.value = "Rate" if is_en else "Oran"
                        elif item.discount_type == DiscountType.AMOUNT:
                            cell.value = "Amount" if is_en else "Miktar"
                        else:
                            cell.value = "—"
                    cell.alignment = Alignment(horizontal="center", vertical="center")

                elif col_key == ColumnKey.DISCOUNT_AMOUNT:
                    if offer.discount_enabled and item.discount_type != DiscountType.NONE:
                        if item.discount_type == DiscountType.RATE:
                            # Oran yüzde olarak (örn: %10 -> 0.10)
                            cell.value = float(item.discount_value) / 100.0
                            cell.number_format = num_format_rate
                        else:
                            cell.value = float(item.discount_value)
                            cell.number_format = num_format_currency
                    else:
                        cell.value = 0
                        cell.number_format = num_format_currency
                    cell.alignment = Alignment(horizontal="right", vertical="center")

                elif col_key == ColumnKey.VAT:
                    if offer.vat_enabled:
                        cell.value = float(offer.vat_rate) / 100.0
                        cell.number_format = num_format_rate
                    cell.alignment = Alignment(horizontal="center", vertical="center")

                elif col_key == ColumnKey.TOTAL_AMOUNT:
                    # Formül: =Miktar * Birim Fiyat
                    if qty_letter and price_letter:
                        cell.value = f"={qty_letter}{current_r}*{price_letter}{current_r}"
                    else:
                        cell.value = float(item.total_amount)
                    cell.number_format = num_format_currency
                    cell.alignment = Alignment(horizontal="right", vertical="center")

                elif col_key == ColumnKey.TOTAL_DISCOUNTED:
                    # Formül: İskontolu Toplam Tutar
                    if offer.discount_enabled and disc_val_letter and total_letter:
                        if item.discount_type == DiscountType.RATE:
                            cell.value = f"={total_letter}{current_r}*(1-{disc_val_letter}{current_r})"
                        elif item.discount_type == DiscountType.AMOUNT:
                            cell.value = f"={total_letter}{current_r}-{disc_val_letter}{current_r}"
                        else:
                            cell.value = f"={total_letter}{current_r}"
                    elif total_letter:
                        cell.value = f"={total_letter}{current_r}"
                    else:
                        cell.value = float(item.total_discounted)
                    cell.number_format = num_format_currency
                    cell.alignment = Alignment(horizontal="right", vertical="center")

                elif col_key == ColumnKey.TOTAL_WITH_VAT:
                    # Formül: KDV Dahil Tutar
                    base_letter = disc_total_letter if (offer.discount_enabled and disc_total_letter) else total_letter
                    if offer.vat_enabled and base_letter:
                        vat_rate_val = float(offer.vat_rate) / 100.0
                        cell.value = f"={base_letter}{current_r}*(1+{vat_rate_val})"
                    elif base_letter:
                        cell.value = f"={base_letter}{current_r}"
                    else:
                        cell.value = float(item.total_with_vat)
                    cell.number_format = num_format_currency
                    cell.alignment = Alignment(horizontal="right", vertical="center")

        end_data_row = start_data_row + len(valid_items) - 1

        # 7. Genel Toplam Özeti (Totals Summary Card)
        summary_start_row = end_data_row + 2
        last_col_idx = len(visible_cols) + 1
        summary_lbl_idx = max(last_col_idx - 1, 2)
        summary_val_idx = last_col_idx

        totals_fill = PatternFill(start_color=self.ACCENT_BG, end_color=self.ACCENT_BG, fill_type="solid")

        # 7.1 Ara Toplam
        subtotal_lbl = "Subtotal:" if is_en else "Ara Toplam:"
        ws.cell(row=summary_start_row, column=summary_lbl_idx, value=subtotal_lbl).font = bold_font
        ws.cell(row=summary_start_row, column=summary_lbl_idx).alignment = Alignment(horizontal="right")
        ws.cell(row=summary_start_row, column=summary_lbl_idx).border = cell_border
        ws.cell(row=summary_start_row, column=summary_lbl_idx).fill = totals_fill

        subtotal_cell = ws.cell(row=summary_start_row, column=summary_val_idx)
        if total_letter:
            subtotal_cell.value = f"=SUM({total_letter}{start_data_row}:{total_letter}{end_data_row})"
        else:
            subtotal_cell.value = sum(float(i.total_amount) for i in valid_items)
        subtotal_cell.font = bold_font
        subtotal_cell.number_format = num_format_currency
        subtotal_cell.alignment = Alignment(horizontal="right")
        subtotal_cell.border = cell_border
        subtotal_cell.fill = totals_fill

        current_summary_r = summary_start_row + 1

        # 7.2 İskonto Satırları (Etkinse)
        if offer.discount_enabled:
            # Toplam İskonto
            disc_lbl = "Total Discount:" if is_en else "Toplam İskonto:"
            ws.cell(row=current_summary_r, column=summary_lbl_idx, value=disc_lbl).font = regular_font
            ws.cell(row=current_summary_r, column=summary_lbl_idx).alignment = Alignment(horizontal="right")
            ws.cell(row=current_summary_r, column=summary_lbl_idx).border = cell_border
            ws.cell(row=current_summary_r, column=summary_lbl_idx).fill = totals_fill

            disc_sum_cell = ws.cell(row=current_summary_r, column=summary_val_idx)
            if total_letter and disc_total_letter:
                disc_sum_cell.value = f"=SUM({total_letter}{start_data_row}:{total_letter}{end_data_row})-SUM({disc_total_letter}{start_data_row}:{disc_total_letter}{end_data_row})"
            else:
                disc_sum_cell.value = float(offer_totals.total_discount)
            disc_sum_cell.font = regular_font
            disc_sum_cell.number_format = num_format_currency
            disc_sum_cell.alignment = Alignment(horizontal="right")
            disc_sum_cell.border = cell_border
            disc_sum_cell.fill = totals_fill
            current_summary_r += 1

            # İskonto Sonrası Ara Toplam
            disc_sub_lbl = "Subtotal After Discount:" if is_en else "İskonto Sonrası Ara Toplam:"
            ws.cell(row=current_summary_r, column=summary_lbl_idx, value=disc_sub_lbl).font = bold_font
            ws.cell(row=current_summary_r, column=summary_lbl_idx).alignment = Alignment(horizontal="right")
            ws.cell(row=current_summary_r, column=summary_lbl_idx).border = cell_border
            ws.cell(row=current_summary_r, column=summary_lbl_idx).fill = totals_fill

            disc_sub_cell = ws.cell(row=current_summary_r, column=summary_val_idx)
            if disc_total_letter:
                disc_sub_cell.value = f"=SUM({disc_total_letter}{start_data_row}:{disc_total_letter}{end_data_row})"
            elif total_letter:
                sub_r = summary_start_row
                disc_r = current_summary_r - 1
                lbl_l = get_column_letter(summary_val_idx)
                disc_sub_cell.value = f"={lbl_l}{sub_r}-{lbl_l}{disc_r}"
            else:
                disc_sub_cell.value = float(offer_totals.discounted_subtotal)
            disc_sub_cell.font = bold_font
            disc_sub_cell.number_format = num_format_currency
            disc_sub_cell.alignment = Alignment(horizontal="right")
            disc_sub_cell.border = cell_border
            disc_sub_cell.fill = totals_fill
            current_summary_r += 1

        # 7.3 KDV Satırları (Etkinse)
        if offer.vat_enabled:
            # Toplam KDV
            vat_rate_str = f"{int(offer.vat_rate)}" if offer.vat_rate % 1 == 0 else f"{offer.vat_rate:g}"
            vat_lbl = f"Total VAT (%{vat_rate_str}):" if is_en else f"Toplam KDV (%{vat_rate_str}):"
            ws.cell(row=current_summary_r, column=summary_lbl_idx, value=vat_lbl).font = regular_font
            ws.cell(row=current_summary_r, column=summary_lbl_idx).alignment = Alignment(horizontal="right")
            ws.cell(row=current_summary_r, column=summary_lbl_idx).border = cell_border
            ws.cell(row=current_summary_r, column=summary_lbl_idx).fill = totals_fill

            vat_sum_cell = ws.cell(row=current_summary_r, column=summary_val_idx)
            val_col_l = get_column_letter(summary_val_idx)
            base_sum_r = current_summary_r - 1 if offer.discount_enabled else summary_start_row
            vat_sum_cell.value = f"={val_col_l}{base_sum_r}*({float(offer.vat_rate)}/100)"
            vat_sum_cell.font = regular_font
            vat_sum_cell.number_format = num_format_currency
            vat_sum_cell.alignment = Alignment(horizontal="right")
            vat_sum_cell.border = cell_border
            vat_sum_cell.fill = totals_fill
            current_summary_r += 1

        # 7.4 Genel Toplam (Büyük / Vurgulu)
        if offer.vat_enabled:
            grand_lbl = "Grand Total (Incl. VAT):" if is_en else "KDV Dâhil Genel Toplam:"
        else:
            grand_lbl = "Grand Total:" if is_en else "Genel Toplam:"

        ws.cell(row=current_summary_r, column=summary_lbl_idx, value=grand_lbl).font = Font(name="Calibri", size=11, bold=True, color=self.PRIMARY_COLOR)
        ws.cell(row=current_summary_r, column=summary_lbl_idx).alignment = Alignment(horizontal="right")
        ws.cell(row=current_summary_r, column=summary_lbl_idx).border = cell_border
        ws.cell(row=current_summary_r, column=summary_lbl_idx).fill = PatternFill(start_color="E0E7FF", end_color="E0E7FF", fill_type="solid")

        grand_val_cell = ws.cell(row=current_summary_r, column=summary_val_idx)
        val_col_l = get_column_letter(summary_val_idx)
        if offer.vat_enabled:
            vat_r = current_summary_r - 1
            base_r = current_summary_r - 2 if offer.discount_enabled else summary_start_row
            grand_val_cell.value = f"={val_col_l}{base_r}+{val_col_l}{vat_r}"
        elif offer.discount_enabled:
            disc_sub_r = current_summary_r - 1
            grand_val_cell.value = f"={val_col_l}{disc_sub_r}"
        else:
            grand_val_cell.value = f"={val_col_l}{summary_start_row}"

        grand_val_cell.font = Font(name="Calibri", size=11, bold=True, color=self.PRIMARY_COLOR)
        grand_val_cell.number_format = num_format_currency
        grand_val_cell.alignment = Alignment(horizontal="right")
        grand_val_cell.border = cell_border
        grand_val_cell.fill = PatternFill(start_color="E0E7FF", end_color="E0E7FF", fill_type="solid")

        # 8. Ticari Koşullar & Ek Alan
        section_start_r = current_summary_r + 2
        active_sections = [s for s in offer.sections if s.enabled and s.text.strip()]

        if active_sections:
            terms_title = "OFFER TERMS AND CONDITIONS" if is_en else "TEKLİF KOŞULLARI VE AÇIKLAMALAR"
            ws.cell(row=section_start_r, column=2, value=terms_title).font = Font(name="Calibri", size=11, bold=True, color=self.PRIMARY_COLOR)
            current_sec_r = section_start_r + 1

            for sec in active_sections:
                sec_text = strip_html_tags(sec.text) if sec.rich_text else sec.text.strip()
                if not sec_text:
                    continue

                if sec.section_key == SectionKey.ADDITIONAL_NOTES:
                    # Kullanıcı şartı: Ek Alan için sabit başlık olmasın, doğrudan kullanıcının girdiği metin yer alsın
                    val_cell = ws.cell(row=current_sec_r, column=2, value=sec_text)
                    val_cell.font = regular_font
                    val_cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                    current_sec_r += 1
                else:
                    lbl_cell = ws.cell(row=current_sec_r, column=2, value=f"{sec.section_key.label(offer.document_language)}:")
                    lbl_cell.font = bold_font
                    lbl_cell.alignment = Alignment(horizontal="left", vertical="top")

                    val_cell = ws.cell(row=current_sec_r, column=3, value=sec_text)
                    val_cell.font = regular_font
                    val_cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                    current_sec_r += 1

        # 9. Otomatik Sütun Genişliği Hesaplama
        for col_idx in range(2, last_col_idx + 1):
            col_letter = get_column_letter(col_idx)
            max_len = 0
            for r in range(table_header_row, end_data_row + 1):
                val = ws.cell(row=r, column=col_idx).value
                if val:
                    str_val = str(val)
                    if not str_val.startswith("="):
                        max_len = max(max_len, len(str_val))
                    else:
                        max_len = max(max_len, 12)
            # Başlık uzunluğunu da değerlendir
            header_val = ws.cell(row=table_header_row, column=col_idx).value
            if header_val:
                max_len = max(max_len, len(str(header_val)))

            # Açıklama sütununa özel ferah genişlik
            if col_idx == 3:  # Genellikle Açıklama sütunu
                ws.column_dimensions[col_letter].width = max(max_len + 4, 30)
            else:
                ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

        # İlk boş marjin sütunu (A)
        ws.column_dimensions["A"].width = 3

        # 10. Dosyayı Kaydet
        wb.save(str(target_path))
        return target_path

    def _get_column_header(self, key: ColumnKey, offer: Offer) -> str:
        """Sütun başlığı metnini teklif para ve miktar birimleriyle zenginleştirir."""
        is_en = offer.document_language == DocumentLanguage.EN
        curr_sym = offer.currency_code.symbol
        unit_val = offer.quantity_unit.label(offer.document_language)
        vat_rate_str = f"{int(offer.vat_rate)}" if offer.vat_rate % 1 == 0 else f"{offer.vat_rate:g}"

        if is_en:
            headers: dict[ColumnKey, str] = {
                ColumnKey.PRODUCT_CODE: "Product Code",
                ColumnKey.DESCRIPTION: "Description",
                ColumnKey.QUANTITY: f"Quantity ({unit_val})",
                ColumnKey.PALLET: "Pallet",
                ColumnKey.UNIT_PRICE: f"Unit Price ({curr_sym})",
                ColumnKey.DISCOUNT_TYPE: "Discount Type",
                ColumnKey.DISCOUNT_AMOUNT: "Discount Value",
                ColumnKey.VAT: f"VAT (%{vat_rate_str})",
                ColumnKey.TOTAL_AMOUNT: f"Total ({curr_sym})",
                ColumnKey.TOTAL_DISCOUNTED: f"Discounted Total ({curr_sym})",
                ColumnKey.TOTAL_WITH_VAT: f"Total Incl. VAT ({curr_sym})",
            }
        else:
            headers = {
                ColumnKey.PRODUCT_CODE: "Ürün Kodu",
                ColumnKey.DESCRIPTION: "Ürün Açıklaması",
                ColumnKey.QUANTITY: f"Miktar ({unit_val})",
                ColumnKey.PALLET: "Palet",
                ColumnKey.UNIT_PRICE: f"Birim Fiyat ({curr_sym})",
                ColumnKey.DISCOUNT_TYPE: "İskonto Türü",
                ColumnKey.DISCOUNT_AMOUNT: "İskonto Değeri",
                ColumnKey.VAT: f"KDV (%{vat_rate_str})",
                ColumnKey.TOTAL_AMOUNT: f"Toplam Tutar ({curr_sym})",
                ColumnKey.TOTAL_DISCOUNTED: f"İskontolu Tutar ({curr_sym})",
                ColumnKey.TOTAL_WITH_VAT: f"KDV Dahil ({curr_sym})",
            }
        return headers.get(key, key.label(offer.document_language))
